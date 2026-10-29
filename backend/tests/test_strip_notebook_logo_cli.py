"""`python -m app.cli strip-notebook-logo`: crop the "Gemini Notebook" mark off content
stored before generation did it (owner decision 2026-10-29). Test DB and temp media only."""

from __future__ import annotations

import io
from typing import Any

import pytest
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import cli
from app.models.artifact_version import ArtifactVersion
from app.models.enums import ArtifactType, MediaFileKind
from app.models.media_file import MediaFile
from app.pipeline import orchestrator
from app.render import notebook_logo as nl
from app.storage.service import Storage
from tests.integration.test_notebooklm_pipeline import (
    _artifact,
    _media_for,
    _MediaNotebookLm,
    _row,
)
from tests.test_notebook_logo import _png, _slide


@pytest.fixture
def queued(monkeypatch: pytest.MonkeyPatch) -> list[str]:
    calls: list[str] = []
    from app.tasks.pdf import build_artifact_pdf

    monkeypatch.setattr(
        build_artifact_pdf,
        "apply_async",
        lambda args, queue: calls.append(f"{queue}:{args[0]}"),
    )
    return calls


def _old_version(db: Session, kind: ArtifactType, media: Any = None) -> ArtifactVersion:
    """A version stored the way generation did before the mark was cropped."""
    name = "slide_deck.json" if kind is ArtifactType.slide_deck else "infographic.json"
    row = _row(name)
    provider = _MediaNotebookLm(row, media or _media_for(kind, row))
    artifact = _artifact(db, kind)
    # a context of its own: undoing it must not undo the temp media storage patches
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(orchestrator, "strip_notebook_logo", lambda blobs: blobs)
        version = orchestrator.run_artifact_generation(db, provider, artifact.id)
    assert version is not None
    return version


def _files(
    db: Session, version: ArtifactVersion
) -> dict[tuple[MediaFileKind, int | None], MediaFile]:
    db.expire_all()
    rows = db.scalars(select(MediaFile).where(MediaFile.artifact_version_id == version.id))
    return {(m.kind, m.page_index): m for m in rows}


def test_a_slide_deck_is_cleaned_in_place(
    db_session: Session, media_storage: Storage, queued: list[str]
) -> None:
    version = _old_version(db_session, ArtifactType.slide_deck)
    before = _files(db_session, version)
    old_keys = {k: m.storage_key for k, m in before.items()}

    results = dict(cli.strip_notebook_logo(db_session))

    [(label, outcome)] = [(k, v) for k, v in results.items() if str(version.id) in k]
    assert "slide_deck r1" in label
    assert outcome.startswith("cleaned: 2/2 image(s) cropped")
    assert "deck PDF rebuilt" in outcome and "2 PPTX picture(s) cropped" in outcome
    assert queued == [f"pdf:{version.id}"]
    after = _files(db_session, version)
    assert set(after) == set(before)
    for key, m in after.items():
        assert m.storage_key != old_keys[key]  # new key (and a new row)
        assert m.id != before[key].id
        assert media_storage.exists(m.storage_key)
        assert not media_storage.exists(old_keys[key])  # old file removed after the commit
        data = media_storage.open(m.storage_key)
        assert m.size_bytes == len(data)
        if m.kind is MediaFileKind.slide_png:
            with Image.open(io.BytesIO(data)) as im:
                assert im.size[1] < 768 and nl.is_marked(im)
        elif m.kind is MediaFileKind.deck_pdf:
            assert all(p.mediabox.height < 768 for p in PdfReader(io.BytesIO(data)).pages)
        else:
            prs = Presentation(io.BytesIO(data))
            assert all(next(iter(s.shapes)).image.size[1] < 768 for s in prs.slides)

    # a second run finds nothing to do
    queued.clear()
    again = dict(cli.strip_notebook_logo(db_session))
    assert again[label].startswith("unchanged: 0/2 image(s) cropped, 2 already clean")
    assert queued == []
    assert {k: m.storage_key for k, m in _files(db_session, version).items()} == {
        k: m.storage_key for k, m in after.items()
    }


def test_an_infographic_is_cleaned(
    db_session: Session, media_storage: Storage, queued: list[str]
) -> None:
    version = _old_version(db_session, ArtifactType.infographic)
    results = dict(cli.strip_notebook_logo(db_session))
    [outcome] = [v for k, v in results.items() if str(version.id) in k]
    assert outcome.startswith("cleaned: 1/1 image(s) cropped")
    [png] = list(_files(db_session, version).values())
    assert png.kind is MediaFileKind.infographic_png
    with Image.open(io.BytesIO(media_storage.open(png.storage_key))) as im:
        assert nl.is_marked(im)
    assert queued == [f"pdf:{version.id}"]


def test_dry_run_writes_nothing(db_session: Session, queued: list[str]) -> None:
    version = _old_version(db_session, ArtifactType.slide_deck)
    before = {k: m.storage_key for k, m in _files(db_session, version).items()}
    results = dict(cli.strip_notebook_logo(db_session, dry_run=True))
    [outcome] = [v for k, v in results.items() if str(version.id) in k]
    assert outcome.startswith("would clean: 2/2 image(s) cropped")
    assert "replaces 4 file(s) (deck_pdf, pptx, slide_png)" in outcome
    assert {k: m.storage_key for k, m in _files(db_session, version).items()} == before
    assert queued == []


def test_images_without_the_mark_are_left_alone(db_session: Session, queued: list[str]) -> None:
    row = _row("infographic.json")
    [url] = list(_media_for(ArtifactType.infographic, row))
    version = _old_version(
        db_session, ArtifactType.infographic, {url: _png(_slide((250, 250, 250)))}
    )
    results = dict(cli.strip_notebook_logo(db_session))
    [outcome] = [v for k, v in results.items() if str(version.id) in k]
    assert outcome == "unchanged: 0/1 image(s) cropped, 1 without the mark (kept)"
    assert queued == []


def test_a_storage_error_keeps_the_old_files(
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    media_storage: Storage,
    queued: list[str],
) -> None:
    version = _old_version(db_session, ArtifactType.slide_deck)
    before = {k: m.storage_key for k, m in _files(db_session, version).items()}
    real_put = media_storage.put
    calls = {"n": 0}

    def flaky_put(key: str, data: bytes, mime: str) -> Any:
        calls["n"] += 1
        if calls["n"] == 3:
            raise OSError("disk full")
        return real_put(key, data, mime)

    monkeypatch.setattr(media_storage, "put", flaky_put)
    results = dict(cli.strip_notebook_logo(db_session))
    [outcome] = [v for k, v in results.items() if str(version.id) in k]
    assert outcome.startswith("failed (OSError: disk full")
    after = {k: m.storage_key for k, m in _files(db_session, version).items()}
    assert after == before and all(media_storage.exists(k) for k in after.values())
    # the two new files written before the error were removed
    leftovers = list(media_storage.path_for(next(iter(before.values()))).parent.iterdir())
    assert len(leftovers) == len(before)
    assert queued == []


def test_command_is_registered() -> None:
    args = cli.build_parser().parse_args(["strip-notebook-logo", "--dry-run"])
    assert args.dry_run and args.func is cli._cmd_strip_notebook_logo
