"""T018 — slide renderer (pptx -> deck pdf -> per-slide png)."""

from __future__ import annotations

import io
import shutil
import subprocess

import pytest
from pptx import Presentation

from app.render import RenderError
from app.render.slides_pptx import build_pptx, ensure_slide_toolchain, render_slides, warm_up
from app.schemas.content import Slide, SlideDeckContent

CONTENT = SlideDeckContent(
    slides=[
        Slide(title="What you'll learn", bullets=["Neurons", "Layers", "Training"], notes="intro"),
        Slide(title="How a neuron works", bullets=["Weighted sum", "Non-linearity"], tags=["core"]),
        Slide(title="Key takeaways", bullets=["Depth composes features"], notes="wrap up"),
    ]
)

_HAS_SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
_HAS_PDFTOPPM = shutil.which("pdftoppm")
_needs_toolchain = pytest.mark.skipif(
    not (_HAS_SOFFICE and _HAS_PDFTOPPM), reason="LibreOffice / poppler not installed"
)


def test_build_pptx_is_deterministic_with_one_slide_per_content_slide() -> None:
    a = build_pptx(CONTENT)
    b = build_pptx(CONTENT)
    # python-pptx embeds no timestamps for this content -> byte-identical
    assert a == b
    prs = Presentation(io.BytesIO(a))
    assert len(prs.slides) == len(CONTENT.slides)
    assert a[:2] == b"PK"  # a zip (Office Open XML)


def test_missing_toolchain_raises_render_error(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.render.slides_pptx.shutil.which", lambda _name: None)
    with pytest.raises(RenderError, match="soffice"):
        ensure_slide_toolchain()
    with pytest.raises(RenderError, match="soffice"):
        render_slides(CONTENT.model_dump())


@_needs_toolchain
def test_render_slides_produces_pptx_pdf_and_one_png_per_slide() -> None:
    rendered = render_slides(CONTENT.model_dump())
    assert rendered.pptx[:2] == b"PK"
    assert rendered.deck_pdf[:5] == b"%PDF-"
    assert len(rendered.slide_png) == len(CONTENT.slides)
    for png in rendered.slide_png:
        assert png[:8] == b"\x89PNG\r\n\x1a\n"
        assert len(png) > 1000

    # page count still matches on a re-run (determinism of the count, per AC)
    again = render_slides(CONTENT.model_dump())
    assert len(again.slide_png) == len(rendered.slide_png)


def test_missing_pdftoppm_raises_render_error(monkeypatch: pytest.MonkeyPatch) -> None:
    def _which(name: str) -> str | None:
        return "/usr/bin/soffice" if name in ("soffice", "libreoffice") else None

    monkeypatch.setattr("app.render.slides_pptx.shutil.which", _which)
    with pytest.raises(RenderError, match="pdftoppm"):
        render_slides(CONTENT.model_dump())


@_needs_toolchain
def test_warm_up_renders_a_throwaway_deck() -> None:
    warm_up()  # just needs to not raise


@_needs_toolchain
def test_soffice_process_failure_raises_render_error(monkeypatch: pytest.MonkeyPatch) -> None:
    def _boom(cmd: list[str], **_kwargs: object) -> subprocess.CompletedProcess[bytes]:
        raise subprocess.CalledProcessError(1, cmd, stderr=b"broken pptx")

    monkeypatch.setattr("app.render.slides_pptx.subprocess.run", _boom)
    with pytest.raises(RenderError, match="broken pptx"):
        render_slides(CONTENT.model_dump())


@_needs_toolchain
def test_soffice_process_timeout_raises_render_error(monkeypatch: pytest.MonkeyPatch) -> None:
    def _hang(cmd: list[str], **_kwargs: object) -> subprocess.CompletedProcess[bytes]:
        raise subprocess.TimeoutExpired(cmd, 180)

    monkeypatch.setattr("app.render.slides_pptx.subprocess.run", _hang)
    with pytest.raises(RenderError, match="timed out"):
        render_slides(CONTENT.model_dump())


@_needs_toolchain
def test_soffice_producing_no_pdf_raises_render_error(monkeypatch: pytest.MonkeyPatch) -> None:
    def _noop(cmd: list[str], **_kwargs: object) -> subprocess.CompletedProcess[bytes]:
        return subprocess.CompletedProcess(cmd, 0, stdout=b"", stderr=b"")

    monkeypatch.setattr("app.render.slides_pptx.subprocess.run", _noop)
    with pytest.raises(RenderError, match="no PDF"):
        render_slides(CONTENT.model_dump())


@_needs_toolchain
def test_render_slides_rejects_a_mismatched_deck(monkeypatch: pytest.MonkeyPatch) -> None:
    # force the PNG split to yield the wrong count
    real_glob = __import__("pathlib").Path.glob

    def fake_glob(self: object, pattern: str):  # type: ignore[no-untyped-def]
        out = list(real_glob(self, pattern))  # type: ignore[arg-type]
        return out[:1] if pattern.startswith("slide-") else out

    monkeypatch.setattr("pathlib.Path.glob", fake_glob)
    with pytest.raises(RenderError, match="page count"):
        render_slides(CONTENT.model_dump())
