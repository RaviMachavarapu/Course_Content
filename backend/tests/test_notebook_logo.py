"""Cropping NotebookLM's "Gemini Notebook" mark off generated images (owner decision
2026-10-29): detection on synthetic images (dark and light backgrounds, mark present or
absent), on corners cut from real NotebookLM output, the deck PDF / PPTX rebuild, and the
pipeline step."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pypdf import PdfReader

from app.models.enums import MediaFileKind
from app.pipeline.renderers import MediaBlob, strip_notebook_logo
from app.render import notebook_logo as nl

FIX = Path(__file__).parent / "fixtures" / "notebook_logo"

DARK = (12, 58, 110)
LIGHT = (250, 250, 250)


def _font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    return ImageFont.load_default(size=size)


def _logo(image: Image.Image, *, scale: float = 1.0) -> tuple[int, int]:
    """Draw a "Gemini Notebook" mark like NotebookLM's in the bottom-right corner; returns
    its (top, bottom) rows."""
    w, h = image.size
    draw = ImageDraw.Draw(image)
    bg = image.getpixel((w - 2, h - 2))
    assert isinstance(bg, tuple)
    ink = (0, 0, 0) if sum(bg[:3]) > 380 else (235, 240, 245)
    font = _font(round(11 * scale))
    text = "Gemini Notebook"
    left, top, right, bottom = draw.textbbox((0, 0), text, font=font)
    x = w - round(11 * scale) - (right - left)
    y = h - round(12 * scale) - (bottom - top)
    draw.text((x - left, y - top), text, fill=ink, font=font)
    return y, y + (bottom - top)


def _slide(bg: tuple[int, int, int], *, size: tuple[int, int] = (1376, 768)) -> Image.Image:
    image = Image.new("RGB", size, bg)
    draw = ImageDraw.Draw(image)
    fg = (0, 0, 0) if sum(bg) > 380 else (230, 240, 255)
    draw.text((80, 60), "A slide title", fill=fg, font=_font(48))
    draw.rectangle((80, 200, 900, 520), outline=fg, width=3)
    draw.text((100, 690), "Body text near the bottom-left", fill=fg, font=_font(22))
    return image


def _png(image: Image.Image) -> bytes:
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def _size(png: bytes) -> tuple[int, int]:
    with Image.open(io.BytesIO(png)) as im:
        return im.size


@pytest.mark.parametrize("bg", [DARK, LIGHT], ids=["dark", "light"])
def test_the_mark_is_cropped_with_the_strip_below_it(bg: tuple[int, int, int]) -> None:
    image = _slide(bg)
    top, _bottom = _logo(image)
    result = nl.strip_logo(_png(image), label="slide")
    assert result.cropped > 0
    assert result.size == (1376, 768 - result.cropped)
    assert top - 3 <= result.size[1] <= top  # cut just above the mark
    assert result.cropped <= 768 * nl.MAX_CROP_SHARE + 1
    with Image.open(io.BytesIO(result.data)) as out:
        assert nl.is_marked(out)
        # nothing of the slide above the strip changed
        assert out.convert("RGB").tobytes() == image.crop((0, 0, 1376, result.size[1])).tobytes()


@pytest.mark.parametrize("bg", [DARK, LIGHT], ids=["dark", "light"])
def test_without_the_mark_the_image_is_unchanged(bg: tuple[int, int, int]) -> None:
    png = _png(_slide(bg))
    result = nl.strip_logo(png)
    assert result.cropped == 0 and result.data == png


@pytest.mark.parametrize(
    ("size", "scale"),
    [((1536, 2752), 2.0), ((2752, 1536), 2.0), ((2048, 2048), 2.0)],
    ids=["portrait", "landscape", "square"],
)
def test_infographic_shapes(size: tuple[int, int], scale: float) -> None:
    image = Image.new("RGB", size, LIGHT)
    ImageDraw.Draw(image).rounded_rectangle(
        (100, 100, size[0] - 100, size[1] - 300), radius=40, fill=(230, 110, 40)
    )
    top, _ = _logo(image, scale=scale)
    result = nl.strip_logo(_png(image))
    assert result.cropped > 0
    assert top - 4 <= result.size[1] <= top


def test_a_cropped_image_is_left_alone_on_a_second_run() -> None:
    image = _slide(DARK)
    _logo(image)
    once = nl.strip_logo(_png(image))
    twice = nl.strip_logo(once.data)
    assert twice.already and twice.cropped == 0 and twice.data == once.data
    # and even without the marker, nothing mark-shaped is left to find
    with Image.open(io.BytesIO(once.data)) as out:
        assert nl.find_logo(out.convert("RGB")) is None


def test_content_in_the_corner_is_not_taken_for_the_mark() -> None:
    image = _slide(DARK)
    draw = ImageDraw.Draw(image)
    # a big block of text filling the bottom-right corner, down to the edge
    for row in range(5):
        draw.text((1050, 640 + 26 * row), "Lorem ipsum dolor sit", fill="white", font=_font(22))
    assert nl.strip_logo(_png(image)).cropped == 0


def test_a_line_across_the_corner_is_not_the_mark() -> None:
    image = _slide(LIGHT)
    ImageDraw.Draw(image).line((0, 755, 1376, 755), fill=(0, 0, 0), width=2)
    assert nl.strip_logo(_png(image)).cropped == 0


def test_the_crop_never_exceeds_five_percent() -> None:
    image = _slide(DARK)
    _logo(image, scale=4.0)  # a mark far taller than NotebookLM's
    result = nl.strip_logo(_png(image))
    assert result.cropped <= 768 * nl.MAX_CROP_SHARE + 1


def test_slide_lines_just_above_the_mark_are_kept() -> None:
    image = _slide(DARK)
    top, _ = _logo(image)
    ImageDraw.Draw(image).line((0, top - 4, 1350, top - 4), fill=(200, 220, 255), width=2)
    result = nl.strip_logo(_png(image))
    assert result.cropped > 0
    assert result.size[1] >= top - 2  # the line (rows top-4..top-3) stays


@pytest.mark.parametrize("bg", [DARK, LIGHT], ids=["dark", "light"])
@pytest.mark.parametrize("gap", [6, 9, 12])
@pytest.mark.parametrize("footer", ["line", "text"])
def test_a_footer_above_the_mark_on_a_plain_background_is_not_a_plate(
    bg: tuple[int, int, int], gap: int, footer: str
) -> None:
    """PR #2 review finding 4: a slide's footer line (or text) 6-12 rows above the mark, in
    its columns, over a plain background: no plate (the rows between are the background),
    so the strip is cropped, not covered."""
    image = _slide(bg)
    top, _ = _logo(image)
    fg = (0, 0, 0) if sum(bg) > 380 else (230, 240, 255)
    draw = ImageDraw.Draw(image)
    if footer == "line":
        draw.line((0, top - gap - 1, 1376, top - gap - 1), fill=fg, width=2)
    else:
        text = "Footer text running across the corner of the slide"
        left, t, right, b = draw.textbbox((0, 0), text, font=_font(14))
        draw.text(
            (1376 - (right - left) - 4, top - gap - (b - t) - t), text, fill=fg, font=_font(14)
        )
    box = nl.find_logo(image)
    assert box is not None
    assert not nl.is_plate(image, box, nl.plate_top(image, box))
    result = nl.strip_logo(_png(image))
    assert not result.covered
    assert result.cropped > 0 and result.size[1] >= top - 3


def test_a_synthetic_plate_over_content_is_covered() -> None:
    """NotebookLM's plate geometry on a synthetic slide: a lighter rounded band from ~9 rows
    above the mark to ~9 below, wider than the mark, over text on the same rows."""
    image = _slide(DARK)
    top, bottom = _logo(image)
    box = nl.find_logo(image)
    assert box is not None
    draw = ImageDraw.Draw(image)
    draw.text((900, top - 2), "Slide text on the same rows", fill=(230, 240, 255), font=_font(16))
    draw.rounded_rectangle(
        (box.left - 8, box.top - 9, min(1375, box.right + 6), box.bottom + 9),
        radius=6,
        fill=(36, 84, 136),
    )
    _logo(image)  # the mark again, on the plate
    box = nl.find_logo(image)
    assert box is not None
    assert nl.is_plate(image, box, nl.plate_top(image, box))
    assert nl.strip_logo(_png(image)).covered


# -- corners cut from real NotebookLM output (module 2, 2026-10-28) ---------------------


def _real(name: str, size: tuple[int, int]) -> Image.Image:
    corner = Image.open(FIX / name).convert("RGB")
    canvas = Image.new("RGB", size, corner.getpixel((0, 0)))  # type: ignore[arg-type]
    canvas.paste(corner, (size[0] - corner.size[0], size[1] - corner.size[1]))
    return canvas


def test_real_slide_mark_on_a_plain_background() -> None:
    image = _real("slide_corner_plain.png", (1376, 768))
    box = nl.find_logo(image)
    assert box == nl.LogoBox(left=1270, top=749, right=1365, bottom=756)
    # the slide frame's bottom line (rows 745-747) is kept
    assert nl.crop_height(image) == 748


def test_real_slide_mark_on_its_plate_is_covered_not_cropped() -> None:
    """Over content NotebookLM puts the mark on a lighter rounded plate (rows 740-765), and
    the slide's text runs on the same rows, so any crop would clip it. Owner decision
    2026-10-29: there, cover just the mark with the plate's colour; the image keeps its
    size and every text line stays whole."""
    image = _real("slide_corner_plate.png", (1376, 768))
    box = nl.find_logo(image)
    assert box is not None and box.top == 749
    assert nl.plate_top(image, box) == 740

    result = nl.strip_logo(_png(image))
    assert result.covered and result.cropped == 0 and result.changed
    assert result.size == (1376, 768)
    with Image.open(io.BytesIO(result.data)) as out:
        out = out.convert("RGB")
        assert nl.find_logo(out) is None  # the mark is gone
        # the mark's box is one flat colour now: the plate's
        colors = {
            out.getpixel((x, y))
            for x in range(box.left, box.right + 1)
            for y in range(box.top, box.bottom + 1)
        }
        assert len(colors) == 1
        # rows above the plate are untouched
        assert all(out.getpixel((x, 700)) == image.getpixel((x, 700)) for x in range(0, 1376, 25))
    # a second run leaves it alone
    assert nl.strip_logo(result.data).already


def test_cover_keeps_an_rgba_image_and_its_transparency() -> None:
    """PR #2 review finding 5: covering the mark keeps the image's mode; transparent areas
    stay transparent and the cover is the plate's colour with its alpha."""
    image = _real("slide_corner_plate.png", (1376, 768)).convert("RGBA")
    alpha = Image.new("L", image.size, 255)
    ImageDraw.Draw(alpha).rectangle((0, 0, 99, 99), fill=0)  # a transparent corner
    ImageDraw.Draw(alpha).rectangle((1200, 735, 1375, 767), fill=230)  # translucent plate
    image.putalpha(alpha)
    box = nl.find_logo(image)
    assert box is not None

    result = nl.strip_logo(_png(image))
    assert result.covered
    with Image.open(io.BytesIO(result.data)) as out:
        assert out.mode == "RGBA"
        assert out.getpixel((10, 10))[3] == 0
        assert out.getpixel((500, 300)) == image.getpixel((500, 300))
        cover = {
            out.getpixel((x, y))
            for x in range(box.left, box.right + 1)
            for y in range(box.top, box.bottom + 1)
        }
        assert len(cover) == 1
        [color] = cover
        assert color[3] == 230
        assert max(abs(color[i] - (36, 84, 136)[i]) for i in range(3)) <= 8


def test_cover_keeps_a_palette_image() -> None:
    image = _real("slide_corner_plate.png", (1376, 768)).convert("P", palette=Image.ADAPTIVE)
    box = nl.find_logo(image)
    assert box is not None
    result = nl.strip_logo(_png(image))
    assert result.covered
    with Image.open(io.BytesIO(result.data)) as out:
        assert out.mode == "P"
        assert out.getpalette() == image.getpalette()
        assert out.getpixel((500, 300)) == image.getpixel((500, 300))
        cover = {
            out.getpixel((x, y))
            for x in range(box.left, box.right + 1)
            for y in range(box.top, box.bottom + 1)
        }
        assert len(cover) == 1


def test_real_slide_mark_on_a_plain_background_is_cropped() -> None:
    result = nl.strip_logo(_png(_real("slide_corner_plain.png", (1376, 768))))
    assert result.cropped == 20 and not result.covered and result.changed
    assert result.size == (1376, 748)


def test_real_infographic_mark() -> None:
    image = _real("infographic_corner.png", (1536, 2752))
    box = nl.find_logo(image)
    assert box is not None and (box.top, box.bottom) == (2724, 2738)
    assert nl.crop_height(image) == 2723


# -- NotebookLM's deck PDF and PPTX (one raster image per page/slide) -------------------


def _slides_with_mark(n: int) -> list[bytes]:
    out = []
    for i in range(n):
        image = _slide(DARK if i % 2 else LIGHT)
        _logo(image)
        out.append(_png(image))
    return out


def _pptx(pngs: list[bytes]) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(16256000), Emu(9144000)
    for png in pngs:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(png), 0, Emu(38100), Emu(16256000), Emu(9067800))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_images_pdf_is_one_lossless_page_per_image() -> None:
    pngs = [_png(_slide(DARK)), _png(_slide(LIGHT, size=(1376, 740)))]
    pdf = nl.images_pdf(pngs)
    reader = PdfReader(io.BytesIO(pdf))
    assert [(p.mediabox.width, p.mediabox.height) for p in reader.pages] == [
        (1376, 768),
        (1376, 740),
    ]
    page = reader.pages[1].images[0].image
    assert page.size == (1376, 740)
    assert page.convert("RGB").tobytes() == _slide(LIGHT, size=(1376, 740)).tobytes()
    assert nl.raster_pdf_pages(pdf) == 2


def test_the_deck_pdf_is_rebuilt_from_the_cropped_slides() -> None:
    original = _slides_with_mark(3)
    deck = nl.images_pdf(original)
    cropped = [nl.strip_logo(p).data for p in original]
    rebuilt = nl.rebuild_deck_pdf(deck, cropped)
    heights = [p.mediabox.height for p in PdfReader(io.BytesIO(rebuilt)).pages]
    assert heights == [_size(c)[1] for c in cropped] and all(h < 768 for h in heights)


def test_a_deck_pdf_that_is_not_one_image_per_slide_is_kept() -> None:
    deck = nl.images_pdf(_slides_with_mark(2))
    assert nl.rebuild_deck_pdf(deck, _slides_with_mark(3)) == deck  # page count differs
    assert nl.raster_pdf_pages(b"not a pdf") is None


def test_pptx_pictures_are_cropped_and_refitted() -> None:
    pptx = _pptx(_slides_with_mark(2))
    out, cropped = nl.strip_pptx(pptx, label="deck")
    assert cropped == 2
    prs = Presentation(io.BytesIO(out))
    for slide in prs.slides:
        [pic] = list(slide.shapes)
        w, h = pic.image.size
        assert h < 768
        with Image.open(io.BytesIO(pic.image.blob)) as im:
            assert nl.is_marked(im)
        # fitted inside the original picture box, centred, aspect kept
        assert pic.width == 16256000
        assert abs(pic.height - round(16256000 * h / w)) <= 1
        assert pic.top >= 38100 and pic.top + pic.height <= 38100 + 9067800 + 1
        # the replaced picture is gone from the slide's relationships
        images = [r for r in slide.part.rels.values() if "image" in r.reltype]
        assert len(images) == 1
    again, n = nl.strip_pptx(out)
    assert n == 0 and again == out


def test_a_pptx_without_the_mark_is_returned_as_is() -> None:
    pptx = _pptx([_png(_slide(DARK))])
    assert nl.strip_pptx(pptx) == (pptx, 0)


# -- the pipeline step on NotebookLM's downloads ----------------------------------------


def test_strip_notebook_logo_on_a_slide_deck() -> None:
    pngs = _slides_with_mark(2)
    blobs = [
        MediaBlob(MediaFileKind.slide_png, p, "image/png", page_index=i) for i, p in enumerate(pngs)
    ]
    blobs += [
        MediaBlob(MediaFileKind.deck_pdf, nl.images_pdf(pngs), "application/pdf"),
        MediaBlob(MediaFileKind.pptx, _pptx(pngs), "application/x"),
    ]
    out = strip_notebook_logo(blobs)
    assert [b.kind for b in out] == [b.kind for b in blobs]
    slides = [b.data for b in out[:2]]
    assert all(_size(s)[1] < 768 for s in slides)
    pdf = PdfReader(io.BytesIO(out[2].data))
    assert [p.mediabox.height for p in pdf.pages] == [_size(s)[1] for s in slides]
    prs = Presentation(io.BytesIO(out[3].data))
    assert all(next(iter(s.shapes)).image.size[1] < 768 for s in prs.slides)


def test_strip_notebook_logo_keeps_what_it_cannot_read() -> None:
    blobs = [
        MediaBlob(MediaFileKind.infographic_png, b"\x89PNG not really", "image/png"),
        MediaBlob(MediaFileKind.pptx, b"not a pptx", "application/x"),
        MediaBlob(MediaFileKind.audio_m4a, b"audio", "audio/mp4"),
    ]
    assert strip_notebook_logo(blobs) == blobs


def test_strip_notebook_logo_without_a_mark_keeps_the_deck() -> None:
    pngs = [_png(_slide(DARK))]
    deck = nl.images_pdf(pngs)
    blobs = [
        MediaBlob(MediaFileKind.slide_png, pngs[0], "image/png", page_index=0),
        MediaBlob(MediaFileKind.deck_pdf, deck, "application/pdf"),
    ]
    assert strip_notebook_logo(blobs) == blobs


def test_the_band_medians_are_computed_once_per_image(monkeypatch: pytest.MonkeyPatch) -> None:
    """PR #2 review finding 8: the row/column medians are computed once and thresholded
    twice (ink, then the faint edge), with the same result."""
    image = _real("slide_corner_plain.png", (1376, 768))
    calls: list[int] = []
    real = nl._band
    monkeypatch.setattr(nl, "_band", lambda *a: calls.append(1) or real(*a))
    assert nl.find_logo(image) == nl.LogoBox(left=1270, top=749, right=1365, bottom=756)
    assert len(calls) == 1
